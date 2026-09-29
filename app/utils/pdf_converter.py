"""
pdf_converter.py — Conversion directe PDF généré → Word / Excel
================================================================
Pas de logique de template. On lit le PDF déjà rempli et on le convertit.

• PDF → Word  : pdf2docx  (rendu visuel fidèle, mise en page conservée)
• PDF → Excel : PyMuPDF   (extraction texte structurée, tableau propre)

Usage dans main.py (aucun changement côté routes) :
    from app.utils.pdf_converter import pdf_to_word, pdf_to_excel

    word_path  = pdf_to_word("fiche_signaletique", data)
    excel_path = pdf_to_excel("questionnaire_inventaire", data)
"""

import os
import re
from typing import Dict, List, Optional

import fitz  # PyMuPDF
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

from app.config import OUTPUT_FOLDER


# ═══════════════════════════════════════════════════════════════════════════
# STYLES EXCEL
# ═══════════════════════════════════════════════════════════════════════════

_HEADER_FILL  = PatternFill("solid", fgColor="1F3864")
_SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
_OUI_FILL     = PatternFill("solid", fgColor="C6EFCE")
_NON_FILL     = PatternFill("solid", fgColor="FFC7CE")
_LABEL_FILL   = PatternFill("solid", fgColor="EBF0FA")
_OUI_FONT     = Font(bold=True, color="006100", size=10)
_NON_FONT     = Font(bold=True, color="9C0006", size=10)
_BORDER       = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


# ═══════════════════════════════════════════════════════════════════════════
# CHERCHER LE PDF GÉNÉRÉ DANS OUTPUT_FOLDER
# ═══════════════════════════════════════════════════════════════════════════

def _find_pdf(doc_type: str, data: Dict) -> Optional[str]:
    """
    Cherche le dernier PDF généré pour ce doc_type dans OUTPUT_FOLDER.
    Priorité : nom exact → dernier fichier avec le préfixe.
    """
    safe = (data.get("raison_sociale") or "entreprise") \
        .replace(" ", "_").replace("/", "-")

    # Nom exact
    exact = os.path.join(OUTPUT_FOLDER, f"{doc_type}_{safe}.pdf")
    if os.path.exists(exact):
        return exact

    # Dernier fichier avec le même préfixe (trié par date desc)
    prefix = f"{doc_type}_"
    candidates = [
        f for f in os.listdir(OUTPUT_FOLDER)
        if f.startswith(prefix) and f.endswith(".pdf")
    ]
    if candidates:
        candidates.sort(
            key=lambda f: os.path.getmtime(os.path.join(OUTPUT_FOLDER, f)),
            reverse=True
        )
        return os.path.join(OUTPUT_FOLDER, candidates[0])

    return None


def _safe_name(data: Dict) -> str:
    return (data.get("raison_sociale") or "entreprise") \
        .replace(" ", "_").replace("/", "-")


# ═══════════════════════════════════════════════════════════════════════════
# PDF → WORD   (conversion visuelle fidèle via pdf2docx)
# ═══════════════════════════════════════════════════════════════════════════

def pdf_to_word(doc_type: str, data: Dict,
                output_path: str = None) -> str:
    """
    Convertit le PDF rempli en .docx en conservant la mise en page.
    Utilise pdf2docx pour un rendu visuel pixel-perfect.
    """
    from pdf2docx import Converter

    pdf_path = _find_pdf(doc_type, data)
    if not pdf_path:
        raise FileNotFoundError(
            f"PDF introuvable pour '{doc_type}' dans {OUTPUT_FOLDER}. "
            "Générez d'abord le PDF."
        )

    if not output_path:
        output_path = os.path.join(
            OUTPUT_FOLDER, f"{doc_type}_{_safe_name(data)}.docx"
        )

    print(f"📄 Word  ← {os.path.basename(pdf_path)}")

    cv = Converter(pdf_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()

    print(f"✅ Word  → {os.path.basename(output_path)}")
    return output_path


# ═══════════════════════════════════════════════════════════════════════════
# PDF → EXCEL  (extraction PyMuPDF → tableau structuré)
# ═══════════════════════════════════════════════════════════════════════════

# Textes à ignorer (headers/footers du template)
_IGNORE_RE = re.compile(
    r"^(cabinet\s+el\s+housny|fait\s+par\s*:|mission\s*:|exercice\s+du\s*:|"
    r"date\s*:|réf\.?\s*:|b_\w+|q_\w+|\d+\s*/\s*\d+|page\s+\d+|kiko|cac|amm_)",
    re.IGNORECASE,
)

# Seuil X pour séparer label (gauche) de valeur (droite)
_SPLIT_X = 310.0


def _read_lines(pdf_path: str) -> List[dict]:
    """
    Lit le PDF et retourne une liste de lignes logiques.
    Chaque ligne = { text, x0, x1, y0, size, bold, page, spans[] }
    """
    doc   = fitz.open(pdf_path)
    buckets: dict = {}

    for pn, page in enumerate(doc):
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block["lines"]:
                for sp in line["spans"]:
                    t = sp.get("text", "").strip()
                    if not t:
                        continue
                    key = (pn, round(sp["bbox"][1] / 2.5) * 2.5)
                    buckets.setdefault(key, []).append({
                        "text" : t,
                        "x0"   : round(sp["bbox"][0], 1),
                        "x1"   : round(sp["bbox"][2], 1),
                        "y0"   : round(sp["bbox"][1], 1),
                        "size" : round(sp.get("size", 9.0), 1),
                        "bold" : bool(
                            "Bold" in sp.get("font", "") or
                            sp.get("flags", 0) & 16
                        ),
                    })
    doc.close()

    lines = []
    for (pn, _), spans in sorted(buckets.items()):
        spans.sort(key=lambda s: s["x0"])
        text = " ".join(s["text"] for s in spans).strip()
        if not text or _IGNORE_RE.match(text):
            continue
        lines.append({
            "text"  : text,
            "x0"    : spans[0]["x0"],
            "x1"    : spans[-1]["x1"],
            "y0"    : spans[0]["y0"],
            "size"  : max(s["size"] for s in spans),
            "bold"  : any(s["bold"] for s in spans),
            "page"  : pn,
            "spans" : spans,
        })
    return lines


def _kind(line: dict) -> str:
    """Classifie une ligne."""
    t  = line["text"].strip()
    tu = t.upper()
    sz = line["size"]
    bld = line["bold"]
    x0  = line["x0"]
    spans = line["spans"]

    if sz > 13 and bld:
        return "title"
    if bld and sz >= 10.5:
        return "section"
    if re.match(r"^\d+[\.\d]*\s+\w", t):
        return "section"
    if tu in {"OUI", "NON", "N.A", "N/A", "FAIT", "OUI/NON/N.A", "OUI/NON"}:
        return "oui_non"
    # Deux spans : un à gauche (<310) et un à droite (>310)
    left  = [s for s in spans if s["x0"] < _SPLIT_X]
    right = [s for s in spans if s["x0"] >= _SPLIT_X]
    if left and right:
        return "label_val"
    if x0 >= _SPLIT_X + 15:
        return "value"
    if "?" in t:
        return "question"
    if t.endswith(":") or (bld and sz < 10):
        return "label"
    return "plain"


def pdf_to_excel(doc_type: str, data: Dict,
                 output_path: str = None) -> str:
    """
    Convertit le PDF rempli en .xlsx structuré.
    • Titres et sections en-tête colorés.
    • Champs d'identification : Label | Valeur sur 2 colonnes.
    • Questions OUI/NON : Question | Réponse (colorée) | Commentaire.
    """
    pdf_path = _find_pdf(doc_type, data)
    if not pdf_path:
        raise FileNotFoundError(
            f"PDF introuvable pour '{doc_type}' dans {OUTPUT_FOLDER}. "
            "Générez d'abord le PDF."
        )

    if not output_path:
        output_path = os.path.join(
            OUTPUT_FOLDER, f"{doc_type}_{_safe_name(data)}.xlsx"
        )

    print(f"📊 Excel ← {os.path.basename(pdf_path)}")

    lines = _read_lines(pdf_path)
    wb    = Workbook()
    ws    = wb.active
    ws.title = _sheet_title(doc_type)
    ws.sheet_view.showGridLines = False

    is_q = "questionnaire" in doc_type or doc_type in (
        "acceptation_mission", "maintien_mission"
    )

    # Largeurs colonnes
    if is_q:
        ws.column_dimensions["A"].width = 62
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 36
    else:
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 48

    xrow = 1
    q_header_written = False
    i = 0

    while i < len(lines):
        line  = lines[i]
        kd    = _kind(line)
        spans = line["spans"]

        # ── Titre ──────────────────────────────────────────────────────────
        if kd == "title":
            ws.row_dimensions[xrow].height = 22
            c = ws.cell(xrow, 1, line["text"])
            c.font      = Font(bold=True, size=14, color="FFFFFF")
            c.fill      = _HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = _BORDER
            ws.merge_cells(xrow, 1, xrow, 3)
            xrow += 2; i += 1; continue

        # ── Section ────────────────────────────────────────────────────────
        if kd == "section":
            ws.row_dimensions[xrow].height = 15
            c = ws.cell(xrow, 1, line["text"])
            c.font      = Font(bold=True, size=11, color="1F3864")
            c.fill      = _SECTION_FILL
            c.alignment = Alignment(vertical="center")
            c.border    = _BORDER
            ws.merge_cells(xrow, 1, xrow, 3)
            xrow += 1; i += 1; continue

        # ── En-tête colonnes questions (une seule fois) ────────────────────
        if kd in ("question", "oui_non") and is_q and not q_header_written:
            q_header_written = True
            for col, label in enumerate(
                    ("Question / Diligence", "Réponse", "Commentaire"), 1):
                c = ws.cell(xrow, col, label)
                c.font      = Font(bold=True, size=10, color="FFFFFF")
                c.fill      = _HEADER_FILL
                c.alignment = Alignment(horizontal="center",
                                        vertical="center", wrap_text=True)
                c.border    = _BORDER
            ws.row_dimensions[xrow].height = 15
            xrow += 1

        # ── Question ───────────────────────────────────────────────────────
        if kd in ("question", "oui_non") and is_q:
            q_txt   = line["text"]
            answer  = ""
            comment = ""

            # Réponse déjà sur la même ligne (spans à droite)
            right = [s for s in spans if s["x0"] >= _SPLIT_X]
            if right:
                answer  = right[0]["text"].strip()
                comment = " ".join(s["text"] for s in right[1:]).strip()

            # Réponse sur la ligne suivante
            if not answer and i + 1 < len(lines):
                nxt = lines[i + 1]["text"].strip().upper()
                if nxt in {"OUI", "NON", "N.A", "N/A", "FAIT"}:
                    answer = lines[i + 1]["text"].strip()
                    i += 1

            ws.row_dimensions[xrow].height = 28

            cq = ws.cell(xrow, 1, q_txt)
            cq.font      = Font(size=9)
            cq.alignment = Alignment(vertical="center", wrap_text=True)
            cq.border    = _BORDER

            ca = ws.cell(xrow, 2, answer)
            au = answer.strip().upper()
            if au in {"OUI", "FAIT"}:
                ca.font = _OUI_FONT; ca.fill = _OUI_FILL
            elif au == "NON":
                ca.font = _NON_FONT; ca.fill = _NON_FILL
            else:
                ca.font = Font(size=9, bold=True)
            ca.alignment = Alignment(horizontal="center", vertical="center")
            ca.border    = _BORDER

            cc = ws.cell(xrow, 3, comment)
            cc.font      = Font(size=8.5, italic=True)
            cc.alignment = Alignment(vertical="center", wrap_text=True)
            cc.border    = _BORDER

            xrow += 1; i += 1; continue

        # ── Paire Label | Valeur ───────────────────────────────────────────
        if kd == "label_val":
            left  = [s for s in spans if s["x0"] < _SPLIT_X]
            right = [s for s in spans if s["x0"] >= _SPLIT_X]
            lbl   = " ".join(s["text"] for s in left).strip().rstrip(":")
            val   = " ".join(s["text"] for s in right).strip()

            ws.row_dimensions[xrow].height = 16
            cl = ws.cell(xrow, 1, lbl)
            cl.font      = Font(bold=True, size=9.5)
            cl.fill      = _LABEL_FILL
            cl.alignment = Alignment(vertical="center")
            cl.border    = _BORDER

            cv = ws.cell(xrow, 2, val)
            cv.font      = Font(size=9.5)
            cv.alignment = Alignment(vertical="center", wrap_text=True)
            cv.border    = _BORDER

            xrow += 1; i += 1; continue

        # ── Label seul ────────────────────────────────────────────────────
        if kd == "label":
            val = ""
            if i + 1 < len(lines) and _kind(lines[i + 1]) == "value":
                val = lines[i + 1]["text"]; i += 1

            ws.row_dimensions[xrow].height = 16
            cl = ws.cell(xrow, 1, line["text"].rstrip(":"))
            cl.font      = Font(bold=True, size=9.5)
            cl.fill      = _LABEL_FILL
            cl.alignment = Alignment(vertical="center")
            cl.border    = _BORDER

            cv = ws.cell(xrow, 2, val)
            cv.font      = Font(size=9.5)
            cv.alignment = Alignment(vertical="center", wrap_text=True)
            cv.border    = _BORDER

            xrow += 1; i += 1; continue

        # ── Texte libre ────────────────────────────────────────────────────
        c = ws.cell(xrow, 1, line["text"])
        c.font      = Font(bold=line["bold"], size=max(8, line["size"] * 0.78))
        c.alignment = Alignment(wrap_text=True)
        ws.merge_cells(xrow, 1, xrow, 3)
        xrow += 1; i += 1

    ws.freeze_panes = "A2"
    wb.save(output_path)
    print(f"✅ Excel → {os.path.basename(output_path)}")
    return output_path


def _sheet_title(doc_type: str) -> str:
    return {
        "fiche_signaletique":         "Fiche signalétique",
        "acceptation_mission":        "Acceptation mission",
        "maintien_mission":           "Maintien mission",
        "questionnaire_pri":          "Guide PRI",
        "questionnaire_inventaire":   "Inventaire",
        "questionnaire_verification": "Vérifications",
        "questionnaire_evenement":    "Evénements",
        "questionnaire_fin":          "Fin de mission",
    }.get(doc_type, "Document")[:31]
